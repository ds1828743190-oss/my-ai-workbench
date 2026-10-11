# -*- coding: utf-8 -*-
"""投资台账模板.xlsx 生成脚本（空模板，重建用）
锚点设计：
  持仓:   表头1 | 示例2-3 | 预留4-13 | 合计14
  交易记录: 表头1 | 示例2-3 | 预留4-13
  净值月度: 表头1 | 示例2 | 预留3-7
  策略卡:  纯文本说明区
"""
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import date

OUT = str(__import__("pathlib").Path(__file__).resolve().parent / "投资台账模板.xlsx")

# ---------- 配色（商务蓝） ----------
def xl_color(css):
    v = css.removeprefix("#").upper()
    assert len(v) == 6
    return "FF" + v

C_HEADER_BG = "4472C4"   # 表头
C_TOTAL_BG  = "2F5597"   # 汇总
C_INPUT_BG  = "D9E2F3"   # 输入浅底
C_CALC_BG   = "F2F2F2"   # 计算浅灰
C_UP        = "C00000"   # 盈利-红（中国习惯红涨）
C_DOWN      = "006100"   # 亏损-绿
C_BORDER    = "BFBFBF"

thin = Side(style="thin", color=C_BORDER)
BORDER_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
F_HEAD = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
F_BODY = Font(name="微软雅黑", size=10)
F_TOTAL = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
F_SEC = Font(name="微软雅黑", bold=True, color="2F5597", size=11)
FILL_HEAD = PatternFill("solid", fgColor=xl_color("#" + C_HEADER_BG))
FILL_INPUT = PatternFill("solid", fgColor=xl_color("#" + C_INPUT_BG))
FILL_CALC = PatternFill("solid", fgColor=xl_color("#" + C_CALC_BG))
FILL_TOTAL = PatternFill("solid", fgColor=xl_color("#" + C_TOTAL_BG))
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_R = Alignment(horizontal="right", vertical="center")
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = openpyxl.Workbook()
wb.properties.title = "投资台账"

def style_header(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths)):
        col = openpyxl.utils.get_column_letter(i + 1)
        c = ws.cell(row=1, column=i + 1, value=h)
        c.font = F_HEAD; c.fill = FILL_HEAD; c.alignment = AL_C; c.border = BORDER_ALL
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

def put(ws, row, col, value, kind="text", fill=None, align=None):
    """kind: text/num3/share/money/pct/date"""
    c = ws.cell(row=row, column=col, value=value)
    c.font = F_BODY; c.border = BORDER_ALL
    fmt = {"num3": "#,##0.000", "share": "#,##0", "money": "¥#,##0.00",
           "pct": "0.00%", "date": "YYYY-MM-DD", "num2": "0.00"}.get(kind)
    if fmt: c.number_format = fmt
    c.alignment = align if align else (AL_L if kind in ("text",) else AL_R)
    if fill: c.fill = fill
    return c

# ============================================================
# Sheet1 持仓   表头1 | 示例2-3 | 预留4-13 | 合计14
# ============================================================
ws = wb.active; ws.title = "持仓"
H = ["代码", "名称", "类型", "买入均价(元)", "持仓份额", "成本金额(元)",
     "最新价(元)", "最新市值(元)", "浮动盈亏(元)", "盈亏率", "备注"]
W = [10, 16, 10, 12, 10, 14, 12, 14, 14, 10, 18]
style_header(ws, H, W)
# 输入列: A B C D E G K | 计算列: F H I J
INPUT_COLS, CALC_COLS = [1, 2, 3, 4, 5, 7, 11], [6, 8, 9, 10]
KIND = {4: "num3", 5: "share", 6: "money", 7: "num3", 8: "money", 9: "money", 10: "pct"}
samples = [
    ["510890", "红利低波ETF", "红利ETF", 1.050, 10000, None, 1.120, None, None, None, "示例行，可删除"],
    ["511010", "国债ETF", "国债ETF", 101.200, 100, None, 101.800, None, None, None, "示例行，可删除"],
]
for r in range(2, 14):  # 数据2-3 + 预留4-13
    sample = samples[r - 2] if r <= 3 else None
    for ci in range(1, 12):
        fill = FILL_INPUT if ci in INPUT_COLS else FILL_CALC
        v = sample[ci - 1] if sample else None
        put(ws, r, ci, v, kind=KIND.get(ci, "text"), fill=fill)
    # 计算列公式（带空值/除零保护）
    ws.cell(row=r, column=6).value = f'=IF(OR(D{r}="",E{r}=""),"",D{r}*E{r})'
    ws.cell(row=r, column=8).value = f'=IF(OR(E{r}="",G{r}=""),"",E{r}*G{r})'
    ws.cell(row=r, column=9).value = f'=IF(OR(F{r}="",H{r}=""),"",H{r}-F{r})'
    ws.cell(row=r, column=10).value = f'=IF(OR(F{r}="",F{r}=0,H{r}=""),"",I{r}/F{r})'
# 合计行14
for ci in range(1, 12):
    c = ws.cell(row=14, column=ci); c.font = F_TOTAL; c.fill = FILL_TOTAL; c.border = BORDER_ALL
    c.alignment = AL_R if ci in (4, 5, 6, 7, 8, 9, 10) else AL_C
    c.number_format = {"6": "¥#,##0.00", "8": "¥#,##0.00", "9": "¥#,##0.00", "10": "0.00%"}.get(str(ci), "General")
ws["B14"] = "合计"
ws["F14"] = "=SUM(F2:F13)"; ws["H14"] = "=SUM(H2:H13)"; ws["I14"] = "=SUM(I2:I13)"
ws["J14"] = '=IF(OR(F14="",F14=0),"",I14/F14)'
# 条件格式：浮盈/盈亏率 >0红 <0绿（中国习惯）
red_f, green_f = Font(color=C_UP), Font(color=C_DOWN)
for rng in ("I2:I13", "J2:J13"):
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=red_f))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=green_f))
dv = DataValidation(type="list", formula1='"红利ETF,国债ETF,股票,可转债,其他"', allow_blank=True)
ws.add_data_validation(dv); dv.add("C2:C13")

# ============================================================
# Sheet2 交易记录   表头1 | 示例2-3 | 预留4-13
# ============================================================
ws2 = wb.create_sheet("交易记录")
H2 = ["日期", "操作", "代码", "名称", "价格(元)", "份额", "成交金额(元)",
      "费用(元)", "实际支出/收入(元)", "策略依据", "备注"]
W2 = [12, 8, 10, 16, 11, 9, 14, 9, 16, 22, 18]
style_header(ws2, H2, W2)
K2 = {1: "date", 5: "num3", 6: "share", 7: "money", 8: "money", 9: "money"}
S2 = [
    [date(2026, 10, 9), "买入", "510890", "红利低波ETF", 1.050, 10000, None, 5.00, None, "建仓：低位高股息介入", "示例行，可删除"],
    [date(2026, 10, 9), "买入", "511010", "国债ETF", 101.200, 100, None, 5.00, None, "打底：国债压舱", "示例行，可删除"],
]
for r in range(2, 14):
    sample = S2[r - 2] if r <= 3 else None
    for ci in range(1, 12):
        fill = FILL_INPUT if ci in (1, 2, 3, 4, 5, 6, 8, 10, 11) else FILL_CALC
        v = sample[ci - 1] if sample else None
        put(ws2, r, ci, v, kind=K2.get(ci, "text"), fill=fill)
    ws2.cell(row=r, column=7).value = f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})'
    ws2.cell(row=r, column=9).value = (f'=IF(G{r}="","",IF(B{r}="买入",'
                                       f'G{r}+IF(H{r}="",0,H{r}),G{r}-IF(H{r}="",0,H{r})))')
ws2.auto_filter.ref = "A1:K13"
dv2 = DataValidation(type="list", formula1='"买入,卖出,分红"', allow_blank=True)
ws2.add_data_validation(dv2); dv2.add("B2:B13")

# ============================================================
# Sheet3 净值月度   表头1 | 示例2 | 预留3-7
# ============================================================
ws3 = wb.create_sheet("净值月度")
H3 = ["年月", "当月投入(元)", "累计投入(元)", "期末总资产(元)", "当月盈亏(元)",
      "累计收益率", "月末股债利差(%)", "操作摘要", "备注"]
W3 = [10, 14, 14, 15, 13, 11, 14, 26, 16]
style_header(ws3, H3, W3)
K3 = {2: "money", 3: "money", 4: "money", 5: "money", 6: "pct", 7: "num2"}
S3 = [["2026-10", 20620, None, 21380, None, None, 5.80, "建仓红利ETF+国债ETF", "示例行，可删除"]]
for r in range(2, 8):  # 示例2 + 预留3-7
    sample = S3[r - 2] if r == 2 else None
    for ci in range(1, 10):
        fill = FILL_INPUT if ci in (1, 2, 4, 7, 8, 9) else FILL_CALC
        v = sample[ci - 1] if sample else None
        put(ws3, r, ci, v, kind=K3.get(ci, "text"), fill=fill)
    if r == 2:
        ws3.cell(row=r, column=3).value = f'=IF(B{r}="","",B{r})'
        ws3.cell(row=r, column=5).value = f'=IF(D{r}="","",D{r}-N(B{r}))'
    else:
        ws3.cell(row=r, column=3).value = f'=IF(AND(B{r}="",C{r-1}=""),"",N(C{r-1})+N(B{r}))'
        ws3.cell(row=r, column=5).value = f'=IF(OR(D{r}="",D{r-1}=""),"",D{r}-D{r-1}-N(B{r}))'
    ws3.cell(row=r, column=6).value = f'=IF(OR(C{r}="",C{r}=0,D{r}=""),"",(D{r}-C{r})/C{r})'
for rng in ("E2:E7", "F2:F7"):
    ws3.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=red_f))
    ws3.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=green_f))

# ============================================================
# Sheet4 策略卡   纯文本说明（无公式）
# ============================================================
ws4 = wb.create_sheet("策略卡")
ws4.column_dimensions["A"].width = 14
ws4.column_dimensions["B"].width = 80
rows = [
    ("title", "股债利差策略卡（规则固化，防情绪化操作）"),
    ("sec", "一、核心信号：沪深300股债利差 = 1/沪深300市盈率 − 10年期国债收益率"),
    ("kv", ("信号区间", "操作动作")),
    ("kv", ("> 6%", "每月定投红利ETF——挑托管费+管理费最低的前五之一，优先每季度或半年分红一次的产品（实打实的现金流）。资产账面被套不慌，持续定投拉低成本")),
    ("kv", ("4% ~ 6%", "持有观望：按原计划定投不停手，但不加仓")),
    ("kv", ("< 4%", "减持一部分权益仓位，换入国债ETF（泡沫区，先落袋）")),
    ("blank", ""),
    ("sec", "二、交叉验证：巴菲特指标 = A股总市值 ÷ 最近12个月GDP"),
    ("kv", ("< 80%", "市场低估——配合股债利差，积极定投")),
    ("kv", ("80% ~ 100%", "中性——按信号执行")),
    ("kv", ("> 100%", "市场偏贵——谨慎加仓，作为减仓的辅助确认信号")),
    ("blank", ""),
    ("sec", "三、执行纪律"),
    ("kv", ("纪律1", "只看数据不看情绪：每次操作前，先看当日简报的利差读数，按区间动作，不自作主张")),
    ("kv", ("纪律2", "每次买卖当天填【交易记录】，写清策略依据——月度复盘用")),
    ("kv", ("纪律3", "每月最后一个交易日：更新【持仓】的最新价，并在【净值月度】记一行")),
    ("kv", ("纪律4", "本金在低位高股息介入，通过定投保证名义价值不亏损，泡沫产生时卖出；老老实实收分红")),
    ("blank", ""),
    ("sec", "四、工作台联动"),
    ("kv", ("每日简报", "AI 每个交易日 17:35 自动拉取沪深300估值与国债收益率 → 计算股债利差 → 判定区间 → 生成简报（WorkBuddy 自动化，无需手动查询）")),
    ("kv", ("数据来源", "大智慧MCP（估值/行情）＋ WorkBuddy 本地金融数据技能")),
    ("kv", ("月度复盘", "每月让 AI 汇总【交易记录】与【净值月度】，输出执行纪律评分与复利进度")),
]
r = 1
for kind, content in rows:
    if kind == "title":
        c = ws4.cell(row=r, column=1, value=content)
        c.font = Font(name="微软雅黑", bold=True, size=14, color="1F3864")
    elif kind == "sec":
        c = ws4.cell(row=r, column=1, value=content)
        c.font = F_SEC
    elif kind == "kv":
        a, b = content
        ca = ws4.cell(row=r, column=1, value=a); cb = ws4.cell(row=r, column=2, value=b)
        ca.font = Font(name="微软雅黑", bold=True, size=10); ca.alignment = AL_C
        cb.font = F_BODY; cb.alignment = AL_L
        ca.border = BORDER_ALL; cb.border = BORDER_ALL
        ca.fill = FILL_INPUT; cb.fill = PatternFill("solid", fgColor=xl_color("#FAFAFA"))
    r += 1

wb.save(OUT)
print("saved:", OUT)
